from __future__ import annotations
from pathlib import Path
from typing import List, Dict, Any
from pydantic import BaseModel, Field

class PublicationReadinessResult(BaseModel):
    status: str = Field(description="pass|block")
    placeholders: List[str] = Field(default_factory=list, min_items=0, max_items=200)
    details: Dict[str, Any] = Field(default_factory=dict)

# Minimal initial placeholder catalogue (expandable)
PLACEHOLDERS = [
    "Month YYYY",
    "INSERT CUSTOMER LOGO",
    "Consultant Only",
    "Remove Before Sending",
    "[AM NAME]",
    "[APPROVER NAME]",
    "[Acme Corp]",
]

def validate_assembled_docx(docx_path: Path) -> PublicationReadinessResult:
    """
    Inspect the assembled DOCX for unresolved placeholders in:
    - paragraphs
    - table cells
    - headers
    - footers

    Returns PublicationReadinessResult(status='block'|'pass', placeholders=list).
    """
    try:
        from docx import Document  # python-docx
    except Exception:
        # If we cannot parse the docx, fail closed to require human review
        return PublicationReadinessResult(status="block", placeholders=["docx_parser_unavailable"])
    found: List[str] = []
    try:
        doc = Document(str(docx_path))

        def _scan_paras(container):
            for p in getattr(container, "paragraphs", []) or []:
                txt = (p.text or "").strip()
                if not txt:
                    continue
                for ph in PLACEHOLDERS:
                    if ph in txt:
                        found.append(ph)

        # Body
        _scan_paras(doc)

        # Sections: headers/footers
        for s in getattr(doc, "sections", []) or []:
            try:
                if hasattr(s, "header") and s.header:
                    _scan_paras(s.header)
            except Exception:
                pass
            try:
                if hasattr(s, "footer") and s.footer:
                    _scan_paras(s.footer)
            except Exception:
                pass

        # Tables
        for tbl in getattr(doc, "tables", []) or []:
            for row in getattr(tbl, "rows", []) or []:
                for cell in getattr(row, "cells", []) or []:
                    txt = (getattr(cell, "text", "") or "").strip()
                    if not txt:
                        continue
                    for ph in PLACEHOLDERS:
                        if ph in txt:
                            found.append(ph)

    except Exception:
        return PublicationReadinessResult(status="block", placeholders=["docx_scan_error"])

    status = "block" if found else "pass"
    return PublicationReadinessResult(status=status, placeholders=sorted(set(found)), details={})