#!/usr/bin/env python3
"""
Utility to create a versioned VCISO Word template with a dedicated partnership placeholder.
This uses python-docx to append a placeholder paragraph to the end of the document.
Usage: python tools/insert_placeholder_to_docx.py input.docx output.docx
"""
import sys
from docx import Document

def insert_placeholder(input_path: str, output_path: str, placeholder: str = "{{ partnership_details }}"):
    doc = Document(input_path)
    doc.add_paragraph()
    para = doc.add_paragraph()
    run = para.add_run(placeholder)
    run.bold = True
    doc.save(output_path)

if __name__ == "__main__":
    # Allow optional third argument with comma-separated placeholders, e.g.:
    # "{{ maturity_gauge }},{{ maturity_score }},{{ maturity_score_category }}"
    if len(sys.argv) not in (3, 4):
        print("Usage: python tools/insert_placeholder_to_docx.py input.docx output.docx [comma_separated_placeholders]")
        sys.exit(1)
    in_path, out_path = sys.argv[1], sys.argv[2]
    if len(sys.argv) == 4 and sys.argv[3].strip():
        placeholders = [p.strip() for p in sys.argv[3].split(',') if p.strip()]
        doc = Document(in_path)
        doc.add_paragraph()
        for ph in placeholders:
            para = doc.add_paragraph()
            run = para.add_run(ph)
            run.bold = True
        doc.save(out_path)
    else:
        insert_placeholder(in_path, out_path)
