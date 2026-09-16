import sys
import os
import types
# Stub streamlit to satisfy config import
st_mod = types.ModuleType("streamlit")
st_mod.secrets = {}
st_mod.session_state = {}
sys.modules["streamlit"] = st_mod
# Ensure project root on sys.path for 'export' import when running from tools/
sys.path.append(os.path.dirname(os.path.abspath(os.path.join(__file__, ".."))))

import io
import json
import pptx
from export import create_tabletop_pptx, validate_tabletop_master_plan

def shape_texts(slide):
    parts = []
    for shp in slide.shapes:
        try:
            if hasattr(shp, "text_frame") and shp.has_text_frame:
                parts.append(shp.text_frame.text or "")
            elif hasattr(shp, "text"):
                parts.append(shp.text or "")
        except Exception:
            continue
    return "\n".join([p for p in parts if p])

def notes_text(slide):
    try:
        ns = slide.notes_slide
        return (ns.notes_text_frame.text or "") if ns else ""
    except Exception:
        return ""

mp = {
    "exercise_title": "Exercise",
    "client_name": "ClientCo",
    "housekeeping_rules": ["No blame", "Phones silent"],
    "objectives": ["Test IR thresholds", "Improve comms"],
    "scope": {"included": ["HQ only"], "excluded": ["Prod excluded"], "assumptions": ["Simulated artefacts only"]},
    "radar_chart_data": {"iam": 2, "endpoint": 2.5, "secops": 1.5, "resilience": 2},
    "improvement_actions": [
        {"id": "ACT-01", "priority": "Critical", "recommendation": "Enforce MFA for all users"},
        {"id": "ACT-02", "priority": "High", "recommendation": "Implement immutable backups with MFA delete"},
        {"id": "ACT-03", "priority": "Medium", "recommendation": "Tune SIEM detections for LAPSUS$ TTPs"},
    ],
    "scenarios": [
        {
            "scenario_title": "Phish leads to endpoint compromise",
            "initial_vector": "Phishing",
            "target_assets": ["Laptops"],
            "injects": [
                {
                    "phase_title": "Detect",
                    "simulated_timestamp": "09:05",
                    "scenario_narrative": "SOC alert of suspicious login from new device.",
                    "expected_mature_response": "Quarantine device and reset credentials.",
                    "decision_threshold": "Declare Major if lateral movement observed.",
                    "technical_indicators": ["Impossible travel alert"],
                    "facilitator_probe_questions": ["What logs do you check first?", "Who do you inform?"],
                    "systems_to_check": ["SIEM", "EDR"],
                    "roles_to_engage": ["IR Lead", "Service Desk"],
                    "runbook_references": ["IR-001 Credential Compromise"],
                    "evidence_hunt": ["Recent logins", "Endpoint process tree"],
                    "knowledge_checks": ["MFA enforcement policy", "Password reset flow"],
                    "timebox_hint": "10m",
                },
                {
                    "phase_title": "Contain",
                    "simulated_timestamp": "09:20",
                    "scenario_narrative": "EDR isolates device; user reports business impact.",
                    "expected_mature_response": "Coordinate containment with minimal disruption.",
                    "decision_threshold": "Escalate if multiple hosts affected.",
                    "technical_indicators": ["Multiple isolation events"],
                    "facilitator_probe_questions": ["What is the rollback plan?", "Who approves isolation?"],
                    "systems_to_check": ["EDR", "Asset CMDB"],
                    "roles_to_engage": ["IT Ops", "Comms"],
                    "runbook_references": ["IR-010 Endpoint Isolation"],
                    "evidence_hunt": ["Isolation logs", "Ticketing records"],
                    "knowledge_checks": ["Change control", "BIA awareness"],
                    "timebox_hint": "8m"
                },
                {
                    "phase_title": "Eradicate/Recover",
                    "simulated_timestamp": "09:40",
                    "scenario_narrative": "Golden image redeploy considered.",
                    "expected_mature_response": "Rebuild and validate before reconnect.",
                    "decision_threshold": "Resume operations post-validation only.",
                    "technical_indicators": ["Hash reputations"],
                    "facilitator_probe_questions": ["How do you validate clean state?", "Who signs off?"],
                    "systems_to_check": ["AV", "Patch mgmt"],
                    "roles_to_engage": ["System Owner", "CISO"],
                    "runbook_references": ["IR-020 Reimage Procedure"],
                    "evidence_hunt": ["AV scan results", "Patch status"],
                    "knowledge_checks": ["Hardening baseline", "Sign-off matrix"],
                    "timebox_hint": "12m"
                }
            ],
        },
        {
            "scenario_title": "Ransomware in file server share",
            "initial_vector": "Drive-by download",
            "target_assets": ["File servers"],
            "injects": [
                {
                    "phase_title": "Detect",
                    "simulated_timestamp": "10:10",
                    "scenario_narrative": "Users see encrypted files and ransom note.",
                    "expected_mature_response": "Disable affected shares; begin triage.",
                    "decision_threshold": "Major if more than one share impacted.",
                    "technical_indicators": [".locked extensions"],
                    "facilitator_probe_questions": ["How fast can you disable shares?", "Who informs the board?"],
                    "systems_to_check": ["Storage mgmt", "SIEM"],
                    "roles_to_engage": ["Storage Admin", "Legal"],
                    "runbook_references": ["IR-030 Ransomware Response"],
                    "evidence_hunt": ["Ransom note hash", "Event logs"],
                    "knowledge_checks": ["Notification thresholds", "Insurer clause"],
                    "timebox_hint": "10m"
                },
                {
                    "phase_title": "Contain",
                    "simulated_timestamp": "10:25",
                    "scenario_narrative": "Containment actions underway.",
                    "expected_mature_response": "Block further spread via network controls.",
                    "decision_threshold": "Escalate to disaster if replication hit.",
                    "technical_indicators": ["SMB brute-force"],
                    "facilitator_probe_questions": ["What network controls can you apply?", "Do you isolate subnets?"],
                    "systems_to_check": ["Firewall", "EDR"],
                    "roles_to_engage": ["Network", "IR Lead"],
                    "runbook_references": ["IR-031 Network Containment"],
                    "evidence_hunt": ["Firewall blocks", "Lateral movement traces"],
                    "knowledge_checks": ["Segmentation policy", "Zero Trust"],
                    "timebox_hint": "8m"
                },
                {
                    "phase_title": "Recover",
                    "simulated_timestamp": "10:45",
                    "scenario_narrative": "Plan recovery from immutable backups.",
                    "expected_mature_response": "Restore from known good, validate integrity.",
                    "decision_threshold": "Declare disaster if RPO/RTO breached.",
                    "technical_indicators": ["Backup job success"],
                    "facilitator_probe_questions": ["What is the clean restore point?", "How verify no reinfection?"],
                    "systems_to_check": ["Backup", "EDR"],
                    "roles_to_engage": ["Backup Admin", "Business Owner"],
                    "runbook_references": ["DR-005 File Restore"],
                    "evidence_hunt": ["Restore logs", "EDR attest"],
                    "knowledge_checks": ["Immutable backups", "IR/DR linkage"],
                    "timebox_hint": "12m"
                }
            ]
        }
    ],
}

ok, defects = validate_tabletop_master_plan(mp)
print("VALIDATION_OK", ok)
print("DEFECTS", defects)
data = create_tabletop_pptx(mp)
print("PPTX_BYTES", len(data))
p = pptx.Presentation(io.BytesIO(data))
print("SLIDES_COUNT", len(p.slides))
dump = []
for i, s in enumerate(p.slides, 1):
    title = s.shapes.title.text if s.shapes.title else ""
    txt = shape_texts(s)
    nts = notes_text(s)
    dump.append({"idx": i, "title": title, "text": txt, "notes": nts})
print(json.dumps(dump, ensure_ascii=False, indent=2))