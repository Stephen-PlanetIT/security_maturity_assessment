import io
import os
import sys
import json
import types

# Stub streamlit to satisfy config import if needed
st = types.ModuleType("streamlit")
st.secrets = {}
st.session_state = {}
sys.modules["streamlit"] = st

import pptx
from export import create_tabletop_pptx

def has_tag(slide, tag: str) -> bool:
    for sh in slide.shapes:
        try:
            if getattr(sh, "name", "") == tag:
                return True
        except Exception:
            pass
    return False

def slide_text(slide) -> str:
    parts = []
    for shp in slide.shapes:
        try:
            if hasattr(shp, "text_frame") and shp.has_text_frame:
                parts.append(shp.text_frame.text or "")
        except Exception:
            continue
    return "\n".join([p for p in parts if p])

def run(profile: str = "standard"):
    # Legacy alias for detail profile (retain), plus new explicit envs for compatibility in tools
    os.environ["TABLETOP_PRESENTATION_PROFILE"] = profile
    os.environ["presentation_detail_profile"] = profile  # alias passthrough
    os.environ["exercise_profile"] = os.environ.get("exercise_profile", "blended")

    # Minimal master plan
    mp = {
        "exercise_title": "Exercise",
        "client_name": "ClientCo",
        "housekeeping_rules": ["No blame", "Phones silent"],
        "objectives": ["Test IR thresholds", "Improve comms"],
        "scope": ["HQ only", "Prod excluded"],
        "radar_chart_data": {"iam": 2, "endpoint": 2.5, "secops": 1.5, "resilience": 2},
        "improvement_actions": [
            {"id": "ACT-01", "priority": "Critical", "recommendation": "Enforce MFA for all users"},
            {"id": "ACT-02", "priority": "High", "recommendation": "Implement immutable backups with MFA delete"},
            {"id": "ACT-03", "priority": "Medium", "recommendation": "Tune SIEM detections"},
        ],
        "scenarios": [
            {
                "scenario_title": "Phish leads to endpoint compromise",
                "initial_vector": "Phishing",
                "target_assets": ["Laptops"],
                "injects": [
                    {
                        "phase_title": "Detect",
                        "simulated_timestamp": "09:05",
                        "scenario_narrative": "SOC alert of suspicious login from new device.",
                        "expected_mature_response": "Quarantine device and reset credentials.",
                        "decision_threshold": "Declare Major if lateral movement observed.",
                        "technical_indicators": ["Impossible travel"],
                        "facilitator_probe_questions": ["What logs do you check first?", "Who do you inform?"],
                        "systems_to_check": ["SIEM", "EDR"],
                        "roles_to_engage": ["IR Lead", "Service Desk"],
                        "runbook_references": ["IR-001 Credential Compromise"],
                        "evidence_hunt": ["Recent logins", "Endpoint process tree"],
                        "knowledge_checks": ["MFA enforcement policy", "Password reset flow"],
                        "timebox_hint": "10m",
                    },
                    {
                        "phase_title": "Contain",
                        "simulated_timestamp": "09:20",
                        "scenario_narrative": "EDR isolates device; user reports business impact.",
                        "expected_mature_response": "Coordinate containment with minimal disruption.",
                        "decision_threshold": "Escalate if multiple hosts affected.",
                        "technical_indicators": ["EDR isolate event"],
                        "facilitator_probe_questions": ["What is the rollback plan?", "Who approves isolation?"],
                        "systems_to_check": ["EDR", "Asset CMDB"],
                        "roles_to_engage": ["IT Ops", "Comms"],
                        "runbook_references": ["IR-010 Endpoint Isolation"],
                        "evidence_hunt": ["Isolation logs", "Ticketing records"],
                        "knowledge_checks": ["Change control", "BIA awareness"],
                        "timebox_hint": "8m",
                    },
                    {
                        "phase_title": "Eradicate",
                        "simulated_timestamp": "09:35",
                        "scenario_narrative": "Golden image redeploy considered.",
                        "expected_mature_response": "Rebuild and validate before reconnect.",
                        "decision_threshold": "Resume operations post-validation only.",
                        "technical_indicators": ["AV clean scan"],
                        "facilitator_probe_questions": ["How do you validate clean state?", "Who signs off?"],
                        "systems_to_check": ["AV", "Patch mgmt"],
                        "roles_to_engage": ["System Owner", "CISO"],
                        "runbook_references": ["IR-020 Reimage Procedure"],
                        "evidence_hunt": ["AV scan results", "Patch status"],
                        "knowledge_checks": ["Hardening baseline", "Sign-off matrix"],
                        "timebox_hint": "12m",
                    }
                ],
            },
            {
                "scenario_title": "Ransomware in file server share",
                "initial_vector": "Drive-by download",
                "target_assets": ["File servers"],
                "injects": [
                    {
                        "phase_title": "Detect",
                        "simulated_timestamp": "10:10",
                        "scenario_narrative": "Users see encrypted files and ransom note.",
                        "expected_mature_response": "Disable affected shares; begin triage.",
                        "decision_threshold": "Major if more than one share impacted.",
                        "technical_indicators": [".locked"],
                        "facilitator_probe_questions": ["How fast can you disable shares?", "Who informs the board?"],
                        "systems_to_check": ["Storage mgmt", "SIEM"],
                        "roles_to_engage": ["Storage Admin", "Legal"],
                        "runbook_references": ["IR-030 Ransomware Response"],
                        "evidence_hunt": ["Ransom note hash", "Event logs"],
                        "knowledge_checks": ["Notification thresholds", "Insurer clause"],
                        "timebox_hint": "10m",
                    },
                    {
                        "phase_title": "Contain",
                        "simulated_timestamp": "10:25",
                        "scenario_narrative": "Containment actions underway.",
                        "expected_mature_response": "Block further spread via network controls.",
                        "decision_threshold": "Escalate to disaster if replication hit.",
                        "technical_indicators": ["SMB block events"],
                        "facilitator_probe_questions": ["What network controls can you apply?", "Do you isolate subnets?"],
                        "systems_to_check": ["Firewall", "EDR"],
                        "roles_to_engage": ["Network", "IR Lead"],
                        "runbook_references": ["IR-031 Network Containment"],
                        "evidence_hunt": ["Firewall blocks", "Lateral movement traces"],
                        "knowledge_checks": ["Segmentation policy", "Zero Trust"],
                        "timebox_hint": "8m",
                    },
                    {
                        "phase_title": "Recover",
                        "simulated_timestamp": "10:40",
                        "scenario_narrative": "Plan recovery from immutable backups.",
                        "expected_mature_response": "Restore from known good, validate integrity.",
                        "decision_threshold": "Declare disaster if RPO/RTO breached.",
                        "technical_indicators": ["Restore test logs"],
                        "facilitator_probe_questions": ["What is the clean restore point?", "How verify no reinfection?"],
                        "systems_to_check": ["Backup", "EDR"],
                        "roles_to_engage": ["Backup Admin", "Business Owner"],
                        "runbook_references": ["DR-005 File Restore"],
                        "evidence_hunt": ["Restore logs", "EDR attest"],
                        "knowledge_checks": ["Immutable backups", "IR/DR linkage"],
                        "timebox_hint": "12m",
                    }
                ],
            },
        ],
    }

    data = create_tabletop_pptx(mp)
    print("PPTX_LEN", len(data))
    if not data:
        print("ERROR: Export returned empty bytes. Strict gating may have failed or template unresolved.")
        sys.exit(2)
    p = pptx.Presentation(io.BytesIO(data))

    first_is_front = has_tag(p.slides[0], "STATIC_01_PLANET_FRONT_PAGE") if len(p.slides) else False
    last_is_wrap = has_tag(p.slides[-1], "STATIC_03_PLANET_WRAP_UP") if len(p.slides) else False
    front_count = sum(1 for s in p.slides if has_tag(s, "STATIC_01_PLANET_FRONT_PAGE"))
    wrap_count = sum(1 for s in p.slides if has_tag(s, "STATIC_03_PLANET_WRAP_UP"))

    print(json.dumps({
        "slides_count": len(p.slides),
        "first_is_front": first_is_front,
        "last_is_wrap": last_is_wrap,
        "front_count": front_count,
        "wrap_count": wrap_count,
        "profile": profile
    }))
    # Debug: print indices of static slides to confirm placement
    front_idxs = []
    wrap_idxs = []
    for idx in range(len(p.slides)):
        try:
            s = p.slides[idx]
            if has_tag(s, "STATIC_01_PLANET_FRONT_PAGE"):
                front_idxs.append(idx)
            if has_tag(s, "STATIC_03_PLANET_WRAP_UP"):
                wrap_idxs.append(idx)
        except Exception:
            continue
    print("STATIC_INDEXES", json.dumps({"front": front_idxs, "wrap": wrap_idxs}))

    # Print first 5 slides overview (title + text)
    for i in range(min(5, len(p.slides))):
        s = p.slides[i]
        title = s.shapes.title.text if s.shapes.title else ""
        txt = slide_text(s)
        print(f"IDX{i+1} TITLE: {title}")
        print(f"IDX{i+1} TEXT: {txt[:500]}")
    return 0

if __name__ == "__main__":
    prof = os.environ.get("TABLETOP_PRESENTATION_PROFILE", "standard").strip().lower()
    rc = run(prof)
    sys.exit(rc)