import sys
import types
# Runtime stub for streamlit to satisfy config.py import without installing full Streamlit
st_mod = types.ModuleType("streamlit")
class DummySecrets(dict):
    def get(self, k, d=None):
        return super().get(k, d)
st_mod.secrets = DummySecrets()
st_mod.session_state = {}
sys.modules["streamlit"] = st_mod

import io
import pptx
from export import create_tabletop_pptx

def slide_text(slide):
    parts = []
    try:
        for shp in slide.shapes:
            try:
                if hasattr(shp, "text_frame") and shp.has_text_frame:
                    parts.append(shp.text_frame.text or "")
                else:
                    parts.append(getattr(shp, "text", "") or "")
            except Exception:
                continue
    except Exception:
        pass
    return "\n".join([p for p in parts if p])

def picture_count(pres):
    cnt = 0
    try:
        for s in pres.slides:
            for shp in s.shapes:
                if hasattr(shp, "image") and shp.image is not None:
                    cnt += 1
    except Exception:
        pass
    return cnt

mp = {
    "exercise_title": "Exercise",
    "client_name": "ClientCo",
    "housekeeping_rules": ["No blame", "Phones silent"],
    "objectives": ["Test IR thresholds", "Improve comms"],
    "scope": ["HQ only", "Prod excluded"],
    "radar_chart_data": {"iam": 2, "endpoint": 2.5, "secops": 1.5, "resilience": 2},
    "improvement_actions": [
        {"id": "ACT-01", "priority": "Critical", "recommendation": "Enforce MFA for all users"},
        {"id": "ACT-02", "priority": "High", "recommendation": "Implement immutable backups with MFA delete"},
        {"id": "ACT-03", "priority": "Medium", "recommendation": "Tune SIEM detections for LAPSUS$ TTPs"},
    ],
    "scenarios": [
        {
            "scenario_title": "Phish leads to endpoint compromise",
            "initial_vector": "Phishing",
            "target_assets": ["Laptops"],
            "injects": [
                {
                    "phase_title": "Detect",
                    "simulated_timestamp": "09:05",
                    "scenario_narrative": "SOC alert of suspicious login from new device.",
                    "expected_mature_response": "Quarantine device and reset credentials.",
                    "decision_threshold": "Declare Major if lateral movement observed.",
                    "technical_indicators": ["Impossible travel alert"],
                    "facilitator_probe_questions": ["What logs do you check first?", "Who do you inform?"],
                    "systems_to_check": ["SIEM", "EDR"],
                    "roles_to_engage": ["IR Lead", "Service Desk"],
                    "runbook_references": ["IR-001 Credential Compromise"],
                    "evidence_hunt": ["Recent logins", "Endpoint process tree"],
                    "knowledge_checks": ["MFA enforcement policy", "Password reset flow"],
                    "timebox_hint": "10m",
                },
                {
                    "phase_title": "Contain",
                    "simulated_timestamp": "09:20",
                    "scenario_narrative": "EDR isolates device; user reports business impact.",
                    "expected_mature_response": "Coordinate containment with minimal disruption.",
                    "decision_threshold": "Escalate if multiple hosts affected.",
                    "technical_indicators": ["Multiple isolation events"],
                    "facilitator_probe_questions": ["What is the rollback plan?", "Who approves isolation?"],
                    "systems_to_check": ["EDR", "Asset CMDB"],
                    "roles_to_engage": ["IT Ops", "Comms"],
                    "runbook_references": ["IR-010 Endpoint Isolation"],
                    "evidence_hunt": ["Isolation logs", "Ticketing records"],
                    "knowledge_checks": ["Change control", "BIA awareness"],
                    "timebox_hint": "8m",
                },
                {
                    "phase_title": "Eradicate/Recover",
                    "simulated_timestamp": "09:40",
                    "scenario_narrative": "Golden image redeploy considered.",
                    "expected_mature_response": "Rebuild and validate before reconnect.",
                    "decision_threshold": "Resume operations post-validation only.",
                    "technical_indicators": ["Hash reputations"],
                    "facilitator_probe_questions": ["How do you validate clean state?", "Who signs off?"],
                    "systems_to_check": ["AV", "Patch mgmt"],
                    "roles_to_engage": ["System Owner", "CISO"],
                    "runbook_references": ["IR-020 Reimage Procedure"],
                    "evidence_hunt": ["AV scan results", "Patch status"],
                    "knowledge_checks": ["Hardening baseline", "Sign-off matrix"],
                    "timebox_hint": "12m",
                },
            ],
        },
        {
            "scenario_title": "Ransomware in file server share",
            "initial_vector": "Drive-by download",
            "target_assets": ["File servers"],
            "injects": [
                {
                    "phase_title": "Detect",
                    "simulated_timestamp": "10:10",
                    "scenario_narrative": "Users see encrypted files and ransom note.",
                    "expected_mature_response": "Disable affected shares; begin triage.",
                    "decision_threshold": "Major if more than one share impacted.",
                    "technical_indicators": [".locked extensions"],
                    "facilitator_probe_questions": ["How fast can you disable shares?", "Who informs the board?"],
                    "systems_to_check": ["Storage mgmt", "SIEM"],
                    "roles_to_engage": ["Storage Admin", "Legal"],
                    "runbook_references": ["IR-030 Ransomware Response"],
                    "evidence_hunt": ["Ransom note hash", "Event logs"],
                    "knowledge_checks": ["Notification thresholds", "Insurer clause"],
                    "timebox_hint": "10m",
                },
                {
                    "phase_title": "Contain",
                    "simulated_timestamp": "10:25",
                    "scenario_narrative": "Containment actions underway.",
                    "expected_mature_response": "Block further spread via network controls.",
                    "decision_threshold": "Escalate to disaster if replication hit.",
                    "technical_indicators": ["SMB brute-force"],
                    "facilitator_probe_questions": ["What network controls can you apply?", "Do you isolate subnets?"],
                    "systems_to_check": ["Firewall", "EDR"],
                    "roles_to_engage": ["Network", "IR Lead"],
                    "runbook_references": ["IR-031 Network Containment"],
                    "evidence_hunt": ["Firewall blocks", "Lateral movement traces"],
                    "knowledge_checks": ["Segmentation policy", "Zero Trust"],
                    "timebox_hint": "8m",
                },
                {
                    "phase_title": "Recover",
                    "simulated_timestamp": "10:45",
                    "scenario_narrative": "Plan recovery from immutable backups.",
                    "expected_mature_response": "Restore from known good, validate integrity.",
                    "decision_threshold": "Declare disaster if RPO/RTO breached.",
                    "technical_indicators": ["Backup job success"],
                    "facilitator_probe_questions": ["What is the clean restore point?", "How verify no reinfection?"],
                    "systems_to_check": ["Backup", "EDR"],
                    "roles_to_engage": ["Backup Admin", "Business Owner"],
                    "runbook_references": ["DR-005 File Restore"],
                    "evidence_hunt": ["Restore logs", "EDR attest"],
                    "knowledge_checks": ["Immutable backups", "IR/DR linkage"],
                    "timebox_hint": "12m",
                },
            ],
        },
    ],
}

data = create_tabletop_pptx(mp)
assert isinstance(data, (bytes, bytearray)) and len(data) > 4000
p = pptx.Presentation(io.BytesIO(data))

# Agenda detection: either explicit 'Agenda' or the presence of Objectives and Scope headings on a slide
agenda_ok = False
for s in p.slides:
    txt = slide_text(s)
    if "Agenda" in txt or ("Objectives:" in txt and "Scope:" in txt):
        agenda_ok = True
        break
assert agenda_ok, "Agenda slide missing"

# Maturity Overview: look for explicit title text OR presence of >=2 pictures (radar + gauge)
maturity_ok = any("Maturity Overview" in slide_text(s) for s in p.slides) or (picture_count(p) >= 2)
assert maturity_ok, "Maturity Overview slide or visuals missing"

# Notes: any inject slide should have WGLL text populated in notes
has_notes = False
for s in p.slides:
    try:
        ns = s.notes_slide
        if ns and "What Good Looks Like" in (ns.notes_text_frame.text or ""):
            has_notes = True
            break
    except Exception:
        pass
assert has_notes, "Facilitator notes not populated"

# Priority Actions: explicit title or bullet text
priority_ok = any("Priority Actions" in slide_text(s) for s in p.slides)
assert priority_ok, "Priority Actions slide missing"

print("ALL_PPTX_VERIFIED")